"""
n8n 워크플로우 자동화 및 문서화 스크립트
Playwright를 사용하여 n8n UI를 자동화하고 스크린샷을 캡처하여 워드 문서 생성
"""

import asyncio
import json
from pathlib import Path
from datetime import datetime
from playwright.async_api import async_playwright, Page
from docx import Document
from docx.shared import Inches
import anthropic
import os


class N8nAutomation:
    def __init__(self, n8n_url: str, api_key: str = None):
        """
        Args:
            n8n_url: n8n 인스턴스 URL (예: http://localhost:5678)
            api_key: Claude API 키 (문서 파싱용)
        """
        self.n8n_url = n8n_url
        self.api_key = api_key
        self.screenshots = []
        self.client = anthropic.Anthropic(api_key=api_key) if api_key else None

    async def parse_instruction_document(self, doc_path: str) -> list:
        """
        AI를 사용하여 n8n 사용법 문서를 파싱하고 단계별 지시사항 추출

        Args:
            doc_path: 텍스트 문서 경로

        Returns:
            단계별 지시사항 리스트
        """
        if not self.client:
            print("⚠️  Claude API 키가 없어 수동 파싱 모드로 진행합니다.")
            return []

        with open(doc_path, 'r', encoding='utf-8') as f:
            content = f.read()

        message = self.client.messages.create(
            model="claude-3-5-sonnet-20241022",
            max_tokens=2000,
            messages=[{
                "role": "user",
                "content": f"""다음 n8n 사용법 문서를 분석하고, 단계별로 수행해야 할 작업을 JSON 배열로 추출해주세요.
각 단계는 다음 형식을 따릅니다:
{{
    "step": 1,
    "action": "노드 추가" | "설정 변경" | "연결" | "실행",
    "description": "상세 설명",
    "node_type": "HTTP Request" | "Code" | 등 (해당되는 경우),
    "config": {{설정 정보}}
}}

문서 내용:
{content}

JSON 배열만 반환해주세요."""
            }]
        )

        try:
            response_text = message.content[0].text
            # JSON 부분만 추출
            start = response_text.find('[')
            end = response_text.rfind(']') + 1
            if start != -1 and end > start:
                steps = json.loads(response_text[start:end])
                return steps
        except Exception as e:
            print(f"❌ 문서 파싱 실패: {e}")
            return []

        return []

    async def login_if_needed(self, page: Page, username: str = None, password: str = None):
        """n8n 로그인 (필요한 경우)"""
        try:
            # 로그인 페이지 확인
            if "login" in page.url or await page.locator('input[type="email"]').count() > 0:
                print("🔐 로그인 중...")
                if username and password:
                    await page.fill('input[type="email"]', username)
                    await page.fill('input[type="password"]', password)
                    await page.click('button[type="submit"]')
                    await page.wait_for_load_state('networkidle')
                    print("✅ 로그인 완료")
        except Exception as e:
            print(f"⚠️  로그인 건너뛰기: {e}")

    async def create_workflow(self, page: Page, steps: list):
        """
        파싱된 단계에 따라 n8n 워크플로우 생성

        Args:
            page: Playwright 페이지 객체
            steps: 단계별 지시사항 리스트
        """
        for step in steps:
            print(f"📌 Step {step['step']}: {step['description']}")

            try:
                if step['action'] == '노드 추가':
                    await self.add_node(page, step.get('node_type', 'Manual'))

                elif step['action'] == '설정 변경':
                    await self.configure_node(page, step.get('config', {}))

                elif step['action'] == '연결':
                    await self.connect_nodes(page, step.get('from'), step.get('to'))

                elif step['action'] == '실행':
                    await self.execute_workflow(page)

                # 각 단계 후 스크린샷
                screenshot_path = f"screenshots/step_{step['step']}.png"
                await page.screenshot(path=screenshot_path, full_page=True)
                self.screenshots.append({
                    'step': step['step'],
                    'description': step['description'],
                    'path': screenshot_path
                })

                await asyncio.sleep(1)  # UI 안정화 대기

            except Exception as e:
                print(f"❌ Step {step['step']} 실패: {e}")
                # 에러 스크린샷
                await page.screenshot(path=f"screenshots/error_step_{step['step']}.png")

    async def add_node(self, page: Page, node_type: str):
        """n8n 워크플로우에 노드 추가"""
        # n8n UI에서 '+' 버튼 클릭 또는 더블클릭으로 노드 추가
        try:
            # 캔버스에 더블클릭하여 노드 추가 메뉴 열기
            canvas = page.locator('canvas, .canvas, [data-test-id="canvas"]').first
            await canvas.dblclick()

            # 노드 검색 및 선택
            search_input = page.locator('input[placeholder*="Search"], input[type="search"]').first
            await search_input.fill(node_type)
            await asyncio.sleep(0.5)

            # 첫 번째 검색 결과 클릭
            await page.locator(f'text={node_type}').first.click()

            print(f"✅ {node_type} 노드 추가됨")

        except Exception as e:
            print(f"⚠️  노드 추가 실패 ({node_type}): {e}")

    async def configure_node(self, page: Page, config: dict):
        """노드 설정 변경"""
        try:
            for key, value in config.items():
                # 설정 필드 찾기 및 입력
                field = page.locator(f'input[name="{key}"], textarea[name="{key}"]').first
                await field.fill(str(value))

            print(f"✅ 노드 설정 완료: {list(config.keys())}")

        except Exception as e:
            print(f"⚠️  노드 설정 실패: {e}")

    async def connect_nodes(self, page: Page, from_node: str, to_node: str):
        """노드 연결"""
        print(f"🔗 {from_node} → {to_node} 연결 중...")
        # 실제 구현은 n8n UI 구조에 따라 다름

    async def execute_workflow(self, page: Page):
        """워크플로우 실행"""
        try:
            execute_btn = page.locator('button:has-text("Execute"), button:has-text("실행")').first
            await execute_btn.click()
            await asyncio.sleep(2)
            print("✅ 워크플로우 실행됨")
        except Exception as e:
            print(f"⚠️  실행 버튼을 찾을 수 없습니다: {e}")

    async def capture_workflow_screenshots(self, page: Page):
        """워크플로우 전체 및 상세 스크린샷 캡처"""
        print("📸 스크린샷 캡처 중...")

        # 전체 워크플로우
        await page.screenshot(path="screenshots/workflow_full.png", full_page=True)

        # 각 노드별 상세 (옵션)
        nodes = await page.locator('.node, [data-node-type]').all()
        for i, node in enumerate(nodes):
            try:
                await node.click()
                await asyncio.sleep(0.5)
                await page.screenshot(path=f"screenshots/node_{i}.png")
            except:
                pass

    def create_word_document(self, output_path: str = "n8n_workflow_guide.docx"):
        """
        캡처한 스크린샷으로 워드 문서 생성

        Args:
            output_path: 출력 워드 문서 경로
        """
        print("📝 워드 문서 생성 중...")

        doc = Document()
        doc.add_heading('n8n 워크플로우 가이드', 0)

        # 생성 정보
        doc.add_paragraph(f'생성 일시: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
        doc.add_paragraph(f'n8n URL: {self.n8n_url}')
        doc.add_page_break()

        # 각 단계별 스크린샷 추가
        for screenshot in self.screenshots:
            doc.add_heading(f'Step {screenshot["step"]}: {screenshot["description"]}', 1)
            doc.add_paragraph(screenshot['description'])

            try:
                # 스크린샷 이미지 추가 (폭 6인치로 조정)
                doc.add_picture(screenshot['path'], width=Inches(6))
            except Exception as e:
                doc.add_paragraph(f'[이미지 로드 실패: {e}]')

            doc.add_page_break()

        # 문서 저장
        doc.save(output_path)
        print(f"✅ 워드 문서 생성 완료: {output_path}")

    async def run(self, instruction_doc: str, username: str = None, password: str = None):
        """
        전체 자동화 프로세스 실행

        Args:
            instruction_doc: n8n 사용법이 담긴 텍스트 문서 경로
            username: n8n 로그인 ID (옵션)
            password: n8n 로그인 비밀번호 (옵션)
        """
        # 스크린샷 디렉토리 생성
        Path("screenshots").mkdir(exist_ok=True)

        print("🚀 n8n 자동화 시작...")

        # 1. 문서 파싱 (AI 사용)
        steps = await self.parse_instruction_document(instruction_doc)
        if not steps:
            print("⚠️  AI 파싱 실패. 수동으로 단계를 정의하세요.")
            # 수동으로 단계 정의 예시
            steps = [
                {"step": 1, "action": "노드 추가", "description": "Manual Trigger 추가", "node_type": "Manual"},
                {"step": 2, "action": "노드 추가", "description": "HTTP Request 추가", "node_type": "HTTP Request"},
            ]

        print(f"📋 총 {len(steps)}개 단계 파싱됨")

        # 2. Playwright로 브라우저 자동화
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=False)  # headless=True로 변경 가능
            context = await browser.new_context(viewport={'width': 1920, 'height': 1080})
            page = await context.new_page()

            try:
                # n8n 접속
                print(f"🌐 {self.n8n_url} 접속 중...")
                await page.goto(self.n8n_url)
                await page.wait_for_load_state('networkidle')

                # 로그인
                await self.login_if_needed(page, username, password)

                # 새 워크플로우 생성 페이지로 이동
                try:
                    new_workflow_btn = page.locator('button:has-text("New"), a[href*="new"]').first
                    await new_workflow_btn.click(timeout=5000)
                except:
                    # 이미 워크플로우 편집 페이지에 있을 수 있음
                    pass

                await page.wait_for_load_state('networkidle')

                # 워크플로우 생성
                await self.create_workflow(page, steps)

                # 최종 스크린샷
                await self.capture_workflow_screenshots(page)

            finally:
                await browser.close()

        # 3. 워드 문서 생성
        self.create_word_document()

        print("✅ 모든 작업 완료!")


# 사용 예시
async def main():
    """메인 실행 함수"""

    # 설정
    N8N_URL = "http://localhost:5678"  # 또는 클라우드 n8n URL
    INSTRUCTION_DOC = "n8n_instructions.txt"  # n8n 사용법 텍스트 문서
    CLAUDE_API_KEY = os.getenv("ANTHROPIC_API_KEY")  # 또는 직접 입력

    # 자동화 실행
    automation = N8nAutomation(
        n8n_url=N8N_URL,
        api_key=CLAUDE_API_KEY
    )

    await automation.run(
        instruction_doc=INSTRUCTION_DOC,
        username="your_email@example.com",  # n8n 로그인 정보 (옵션)
        password="your_password"
    )


if __name__ == "__main__":
    asyncio.run(main())
